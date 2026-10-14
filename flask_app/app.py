import csv
import io
import json
import os
import zipfile
from datetime import date, timedelta
from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, make_response, session
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import func, or_

app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "change-this-development-key")
app.config["SQLALCHEMY_DATABASE_URI"] = os.environ.get(
    "DATABASE_URL", "sqlite:///" + os.path.join(app.instance_path if hasattr(app, "instance_path") else os.path.dirname(__file__), "smartcourse.db")
)
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(hours=8)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
os.makedirs(os.path.dirname(app.config["SQLALCHEMY_DATABASE_URI"].replace("sqlite:///", "")), exist_ok=True) if app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite:///") else None
db = SQLAlchemy(app)

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), nullable=False, unique=True, index=True)
    email = db.Column(db.String(255), nullable=False, unique=True, index=True)
    password_hash = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, server_default=func.now())

    def set_password(self, password):
        self.password_hash = generate_password_hash(password)

    def check_password(self, password):
        return check_password_hash(self.password_hash, password)

class Subject(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(160), nullable=False)
    code = db.Column(db.String(40), nullable=False, unique=True)
    department = db.Column(db.String(120), default="General")
    semester = db.Column(db.String(80), default="Semester 1")
    outcomes = db.Column(db.Text, default="")
    units = db.Column(db.Text, default="")
    calendar = db.Column(db.Text, default="")
    resources = db.Column(db.Text, default="")
    website_overview = db.Column(db.Text, default="")
    announcements = db.Column(db.Text, default="")
    case_topic = db.Column(db.String(200), default="")
    case_scenario = db.Column(db.Text, default="")
    case_questions = db.Column(db.Text, default="")
    interactive_prompt = db.Column(db.Text, default="")
    exam_title = db.Column(db.String(200), default="Internal Assessment 1")
    exam_marks = db.Column(db.Integer, default=30)
    questions = db.Column(db.Text, default="")
    answer_key = db.Column(db.Text, default="")
    policy = db.Column(db.Text, default="Disclose AI assistance. Verify claims and citations. Protect private information.")
    improvement_plan = db.Column(db.Text, default="")
    # Per-subject syllabus, program outcomes, CO-PO mapping and Bloom question bank.
    syllabus_text = db.Column(db.Text, default="")
    program_outcomes = db.Column(db.Text, default="")
    co_po_mapping = db.Column(db.Text, default="")
    bloom_questions = db.Column(db.Text, default="[]")
    created_at = db.Column(db.DateTime, server_default=func.now())

class Attendance(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject_id = db.Column(db.Integer, db.ForeignKey("subject.id"), nullable=False)
    student_id = db.Column(db.String(80), nullable=False)
    attendance_date = db.Column(db.Date, nullable=False, default=date.today)
    status = db.Column(db.String(20), nullable=False, default="Present")
    subject = db.relationship("Subject", backref=db.backref("attendance_records", lazy=True, cascade="all, delete-orphan"))

class Mark(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject_id = db.Column(db.Integer, db.ForeignKey("subject.id"), nullable=False)
    student_id = db.Column(db.String(80), nullable=False)
    assessment = db.Column(db.String(160), nullable=False)
    obtained = db.Column(db.Float, nullable=False)
    maximum = db.Column(db.Float, nullable=False)
    subject = db.relationship("Subject", backref=db.backref("mark_records", lazy=True, cascade="all, delete-orphan"))

class ResearchPaper(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject_id = db.Column(db.Integer, db.ForeignKey("subject.id"), nullable=False)
    title = db.Column(db.String(300), nullable=False)
    authors = db.Column(db.String(300), default="")
    year = db.Column(db.Integer, nullable=True)
    doi_url = db.Column(db.String(500), default="")
    summary = db.Column(db.Text, default="")
    verified = db.Column(db.Boolean, default=False)
    subject = db.relationship("Subject", backref=db.backref("research_papers", lazy=True, cascade="all, delete-orphan"))

class QuizQuestion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    subject_id = db.Column(db.Integer, db.ForeignKey("subject.id"), nullable=False)
    question = db.Column(db.String(500), nullable=False)
    answer = db.Column(db.String(500), nullable=False)
    subject = db.relationship("Subject", backref=db.backref("quiz_questions", lazy=True, cascade="all, delete-orphan"))

def seed_data():
    if Subject.query.count():
        return
    samples = [
        dict(name="Computer Science", code="CS101", department="Computer Science", semester="Semester 1",
             outcomes="CO1: Explain computing fundamentals (Understand)\nCO2: Apply algorithms to solve problems (Apply)\nCO3: Evaluate program solutions (Evaluate)",
             units="Programming foundations; Data types and control flow; Algorithms; Testing and debugging; Mini project",
             calendar="Week 1: Introduction\nWeek 2: Core concepts\nWeek 3: Guided practice\nWeek 4: Application\nWeek 5: Review",
             resources="Lecture notes; textbook chapter; practice worksheet"),
        dict(name="Artificial Intelligence", code="AI201", department="Computing", semester="Semester 3",
             outcomes="CO1: Explain core AI concepts (Understand)\nCO2: Apply search and learning methods (Apply)\nCO3: Evaluate AI applications responsibly (Evaluate)",
             units="AI foundations; Search; Machine learning; Neural networks; Responsible AI",
             resources="Lecture slides; verified research articles; practice exercises"),
        dict(name="Mathematics", code="MA101", department="Science", semester="Semester 1",
             outcomes="CO1: Explain mathematical concepts (Understand)\nCO2: Solve structured problems (Apply)\nCO3: Interpret results (Analyze)",
             units="Functions; Algebra; Calculus; Probability; Applications",
             resources="Textbook exercises; worksheets; graphing activities"),
        dict(name="Physics", code="PH101", department="Science", semester="Semester 1",
             outcomes="CO1: Describe physical principles (Understand)\nCO2: Apply equations to scenarios (Apply)\nCO3: Analyze experimental observations (Analyze)",
             units="Measurement; Motion; Forces; Energy; Waves"),
        dict(name="Commerce", code="CM101", department="Commerce", semester="Semester 1",
             outcomes="CO1: Explain accounting concepts (Understand)\nCO2: Prepare basic statements (Apply)\nCO3: Analyze business transactions (Analyze)",
             units="Accounting principles; Journal and ledger; Trial balance; Financial statements; Business analysis")
    ]
    for s in samples:
        subject_data = {
            **s,
            "calendar": s.get("calendar") or "Week 1: Introduction\nWeek 2: Core concepts",
            "website_overview": f"{s['name']} course overview. Add the approved syllabus and course details.",
            "announcements": "Welcome. Check the weekly schedule and assessment deadlines.",
            "case_topic": s.get("units", "").split(";")[0],
            "case_scenario": f"Students apply a core concept from {s['name']} to a realistic problem.",
            "case_questions": "1. What is the main problem?\n2. Which concepts apply?\n3. What alternatives exist?\n4. How would you evaluate the solution?",
            "interactive_prompt": "Explain one key concept and provide a practical example.",
            "syllabus_text": s.get("units", ""),
        }
        db.session.add(Subject(**subject_data))
    db.session.commit()

def migrate_subject_columns():
    """Add columns to existing SQLite installs without deleting existing course data."""
    if not app.config["SQLALCHEMY_DATABASE_URI"].startswith("sqlite:///"):
        return
    from sqlalchemy import inspect, text
    inspector = inspect(db.engine)
    if "subject" not in inspector.get_table_names():
        return
    existing = {column["name"] for column in inspector.get_columns("subject")}
    additions = {
        "syllabus_text": "TEXT DEFAULT ''",
        "program_outcomes": "TEXT DEFAULT ''",
        "co_po_mapping": "TEXT DEFAULT ''",
        "bloom_questions": "TEXT DEFAULT '[]'",
    }
    with db.engine.begin() as connection:
        for column, definition in additions.items():
            if column not in existing:
                connection.execute(text(f"ALTER TABLE subject ADD COLUMN {column} {definition}"))


with app.app_context():
    db.create_all()
    migrate_subject_columns()
    seed_data()

def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if not session.get("user_id"):
            return redirect(url_for("login", next=request.full_path if request.query_string else request.path))
        return view(*args, **kwargs)
    return wrapped_view

@app.before_request
def require_login():
    allowed_endpoints = {"login", "register", "static"}
    if request.endpoint and request.endpoint not in allowed_endpoints and not session.get("user_id"):
        return redirect(url_for("login", next=request.full_path if request.query_string else request.path))

@app.context_processor
def inject_user():
    user = db.session.get(User, session.get("user_id")) if session.get("user_id") else None
    return {"current_user": user}

@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        identity = request.form.get("identity", "").strip()
        password = request.form.get("password", "")
        user = User.query.filter(or_(User.username.ilike(identity), User.email.ilike(identity))).first() if identity else None
        if user and user.check_password(password):
            session.clear()
            session["user_id"] = user.id
            session.permanent = True
            flash("Welcome back, " + user.username + "!", "success")
            next_url = request.form.get("next", "")
            if next_url.startswith("/") and not next_url.startswith("//"):
                return redirect(next_url)
            return redirect(url_for("dashboard"))
        flash("Login failed. Check your username/email and password.", "error")
    return render_template("login.html", next_url=request.args.get("next", ""))

@app.route("/register", methods=["GET", "POST"])
def register():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")
        errors = []
        if len(username) < 3 or len(username) > 80:
            errors.append("Username must be between 3 and 80 characters.")
        if "@" not in email or len(email) > 255:
            errors.append("Enter a valid email address.")
        if len(password) < 8:
            errors.append("Password must contain at least 8 characters.")
        if password != confirm:
            errors.append("Passwords do not match.")
        if User.query.filter(or_(User.username.ilike(username), User.email.ilike(email))).first():
            errors.append("That username or email is already registered.")
        if errors:
            for error in errors:
                flash(error, "error")
            return render_template("register.html")
        user = User(username=username, email=email)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        session.clear()
        session["user_id"] = user.id
        session.permanent = True
        flash("Your account is ready. Welcome to EduOutcome AI!", "success")
        return redirect(url_for("dashboard"))
    return render_template("register.html")

@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    flash("You have been signed out.", "success")
    return redirect(url_for("login"))

def get_subject():
    sid = request.args.get("subject_id", type=int) or request.form.get("subject_id", type=int)
    if sid:
        return db.session.get(Subject, sid) or Subject.query.order_by(Subject.id).first()
    return Subject.query.order_by(Subject.id).first()

@app.context_processor
def inject_globals():
    return {"all_subjects": Subject.query.order_by(Subject.name).all() if app.app_context else [], "active_subject_id": request.args.get("subject_id", type=int) or request.form.get("subject_id", type=int)}

@app.route("/")
def dashboard():
    # Always show sign-in first for visitors without an authenticated session.
    if not session.get("user_id"):
        return redirect(url_for("login", next=request.path))
    s = get_subject()
    if not s:
        return render_template("base.html", page_title="EduOutcome AI", body="<p>Create your first subject.</p>")
    total_subjects = Subject.query.count()
    attendance_count = Attendance.query.filter_by(subject_id=s.id).count()
    marks = Mark.query.filter_by(subject_id=s.id).all()
    avg = (sum((m.obtained / m.maximum) * 100 for m in marks if m.maximum) / len([m for m in marks if m.maximum])) if any(m.maximum for m in marks) else None
    return render_template("dashboard.html", s=s, total_subjects=total_subjects, attendance_count=attendance_count, marks=marks, avg=avg)

@app.route("/subjects/new", methods=["GET", "POST"])
def subject_new():
    if request.method == "POST":
        name = request.form.get("name","").strip()
        code = request.form.get("code","").strip().upper()
        if not name or not code:
            flash("Subject name and course code are required.", "error")
        elif Subject.query.filter_by(code=code).first():
            flash("That course code already exists.", "error")
        else:
            s = Subject(name=name, code=code, department=request.form.get("department","General"),
                        semester=request.form.get("semester","Semester 1"), outcomes=request.form.get("outcomes",""),
                        units=request.form.get("units",""), calendar=request.form.get("calendar",""),
                        resources=request.form.get("resources",""), website_overview=request.form.get("website_overview",""),
                        syllabus_text=request.form.get("syllabus_text", request.form.get("units", "")))
            db.session.add(s); db.session.commit()
            flash("Subject created.", "success")
            return redirect(url_for("planning", subject_id=s.id))
        return render_template("subject_form.html", s=None)
    return render_template("subject_form.html", s=None)

@app.route("/planning", methods=["GET", "POST"])
def planning():
    s = get_subject()
    if request.method == "POST":
        for key in ["name","department","semester","outcomes","units","calendar","resources"]:
            setattr(s, key, request.form.get(key, getattr(s,key) or ""))
        s.syllabus_text = request.form.get("syllabus_text", s.units or "")
        code = request.form.get("code", s.code).strip().upper()
        if code != s.code and Subject.query.filter_by(code=code).first():
            flash("Course code already exists; existing code retained.", "error")
        else:
            s.code = code
        db.session.commit(); flash("Course plan saved.", "success")
        return redirect(url_for("planning", subject_id=s.id))
    return render_template("planning.html", s=s)

@app.route("/website", methods=["GET", "POST"])
def website():
    s = get_subject()
    if request.method == "POST":
        s.website_overview = request.form.get("website_overview","")
        s.announcements = request.form.get("announcements","")
        db.session.commit(); flash("Website details saved.", "success")
        return redirect(url_for("website", subject_id=s.id))
    return render_template("website.html", s=s)

@app.route("/website/export")
def website_export():
    s = get_subject()
    return render_template("public_course.html", s=s)

@app.route("/experiential", methods=["GET", "POST"])
def experiential():
    s = get_subject()
    if request.method == "POST":
        s.case_topic = request.form.get("case_topic","")
        s.case_scenario = request.form.get("case_scenario","")
        s.case_questions = request.form.get("case_questions","")
        db.session.commit(); flash("Case study saved.", "success")
        return redirect(url_for("experiential", subject_id=s.id))
    return render_template("experiential.html", s=s)

@app.route("/quiz/create", methods=["POST"])
def quiz_create():
    s = get_subject()
    questions = [
        (f"What should a learner do first when solving a problem in {s.name}?", "Define and understand the problem."),
        ("Why should a proposed solution be evaluated?", "To check correctness, limitations and suitability."),
        ("What should you do with uncertain AI-generated information?", "Verify it with reliable sources.")
    ]
    for q,a in questions:
        db.session.add(QuizQuestion(subject_id=s.id, question=q, answer=a))
    db.session.commit()
    flash("Three sample quiz questions added.", "success")
    return redirect(url_for("experiential", subject_id=s.id))

@app.route("/interactive", methods=["GET", "POST"])
def interactive():
    s = get_subject()
    if request.method == "POST":
        s.interactive_prompt = request.form.get("interactive_prompt","")
        db.session.commit(); flash("Interactive activity saved.", "success")
        return redirect(url_for("interactive", subject_id=s.id))
    return render_template("interactive.html", s=s)

@app.route("/assessment", methods=["GET", "POST"])
def assessment():
    s = get_subject()
    if request.method == "POST":
        s.exam_title = request.form.get("exam_title","Internal Assessment")
        s.exam_marks = request.form.get("exam_marks",30,type=int)
        s.questions = request.form.get("questions","")
        s.answer_key = request.form.get("answer_key","")
        db.session.commit(); flash("Assessment saved.", "success")
        return redirect(url_for("assessment", subject_id=s.id))
    question_count = len([x for x in (s.questions or "").splitlines() if x.strip()])
    return render_template("assessment.html", s=s, question_count=question_count)

@app.route("/records", methods=["GET", "POST"])
def records():
    s = get_subject()
    if request.method == "POST":
        kind = request.form.get("kind")
        if kind == "attendance":
            student_id = request.form.get("student_id","").strip()
            status = request.form.get("status","Present")
            try: d = date.fromisoformat(request.form.get("attendance_date",""))
            except ValueError: d = date.today()
            if student_id:
                db.session.add(Attendance(subject_id=s.id, student_id=student_id, attendance_date=d, status=status))
                db.session.commit(); flash("Attendance saved.", "success")
        elif kind == "mark":
            student_id = request.form.get("student_id","").strip()
            assessment = request.form.get("assessment","").strip()
            try:
                obtained = float(request.form.get("obtained","0")); maximum = float(request.form.get("maximum","0"))
            except ValueError:
                obtained, maximum = -1, 0
            if student_id and assessment and maximum > 0 and 0 <= obtained <= maximum:
                db.session.add(Mark(subject_id=s.id, student_id=student_id, assessment=assessment, obtained=obtained, maximum=maximum))
                db.session.commit(); flash("Marks saved.", "success")
            else: flash("Enter valid student, assessment, and marks (0 ≤ obtained ≤ maximum).", "error")
        return redirect(url_for("records", subject_id=s.id))
    attendance = Attendance.query.filter_by(subject_id=s.id).order_by(Attendance.attendance_date.desc()).all()
    marks = Mark.query.filter_by(subject_id=s.id).order_by(Mark.id.desc()).all()
    return render_template("records.html", s=s, attendance=attendance, marks=marks)

@app.route("/records/marks.csv")
def marks_csv():
    s = get_subject()
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["student_id","assessment","obtained","maximum","percentage"])
    for m in Mark.query.filter_by(subject_id=s.id).all():
        writer.writerow([m.student_id,m.assessment,m.obtained,m.maximum,round(m.obtained/m.maximum*100,2) if m.maximum else ""])
    response = make_response(output.getvalue())
    response.headers["Content-Disposition"] = f'attachment; filename="{s.code}_marks.csv"'
    response.headers["Content-Type"] = "text/csv; charset=utf-8"
    return response

@app.route("/obe", methods=["GET", "POST"])
def obe():
    s = get_subject()
    if request.method == "POST":
        s.improvement_plan = request.form.get("improvement_plan","")
        db.session.commit(); flash("Improvement plan saved.", "success")
        return redirect(url_for("obe", subject_id=s.id))
    marks = Mark.query.filter_by(subject_id=s.id).all()
    valid = [m for m in marks if m.maximum > 0]
    avg = sum(m.obtained/m.maximum*100 for m in valid)/len(valid) if valid else None
    return render_template("obe.html", s=s, marks=marks, avg=avg)

@app.route("/agent", methods=["GET", "POST"])
def agent():
    s = get_subject()
    answer = None
    question = ""
    if request.method == "POST":
        question = request.form.get("question","").strip().lower()
        if "outcome" in question or "co" in question: answer = s.outcomes or "No course outcomes saved yet."
        elif "unit" in question or "syllabus" in question: answer = s.units or "No syllabus units saved yet."
        elif "resource" in question or "material" in question: answer = s.resources or "No resources saved yet."
        elif "semester" in question or "course" in question: answer = f"{s.name} ({s.code}), {s.semester}, Department: {s.department}."
        else: answer = "I could not find a verified answer in the saved course profile. Please check the syllabus or ask the course coordinator."
    return render_template("agent.html", s=s, answer=answer, question=question)

@app.route("/research", methods=["GET", "POST"])
def research():
    s = get_subject()
    if request.method == "POST":
        title = request.form.get("title","").strip()
        if title:
            year = request.form.get("year","").strip()
            db.session.add(ResearchPaper(subject_id=s.id, title=title, authors=request.form.get("authors",""),
                year=int(year) if year.isdigit() else None, doi_url=request.form.get("doi_url",""),
                summary=request.form.get("summary",""), verified=bool(request.form.get("verified"))))
            db.session.commit(); flash("Research record saved. Verify source metadata before citing.", "success")
        return redirect(url_for("research", subject_id=s.id))
    papers = ResearchPaper.query.filter_by(subject_id=s.id).order_by(ResearchPaper.year.desc()).all()
    return render_template("research.html", s=s, papers=papers)

@app.route("/responsible", methods=["GET", "POST"])
def responsible():
    s = get_subject()
    if request.method == "POST":
        s.policy = request.form.get("policy","")
        db.session.commit(); flash("Responsible AI policy saved.", "success")
        return redirect(url_for("responsible", subject_id=s.id))
    return render_template("responsible.html", s=s)

@app.route("/export/json")
def export_json():
    s = get_subject()
    data = {
        "subject": {c.name: getattr(s,c.name) for c in s.__table__.columns},
        "attendance": [{"student_id":r.student_id,"date":r.attendance_date.isoformat(),"status":r.status} for r in s.attendance_records],
        "marks": [{"student_id":r.student_id,"assessment":r.assessment,"obtained":r.obtained,"maximum":r.maximum} for r in s.mark_records],
        "research_papers": [{"title":r.title,"authors":r.authors,"year":r.year,"doi_url":r.doi_url,"summary":r.summary,"verified":r.verified} for r in s.research_papers],
        "quiz": [{"question":q.question,"answer":q.answer} for q in s.quiz_questions]
    }
    return app.response_class(json.dumps(data, default=str, indent=2), mimetype="application/json",
                              headers={"Content-Disposition": f'attachment; filename="{s.code}_course_data.json"'})


BLOOM_LEVELS = {
    "Remember": ("define", "list", "identify", "recall", "name"),
    "Understand": ("explain", "summarize", "classify", "describe", "interpret"),
    "Apply": ("apply", "demonstrate", "solve", "use", "implement"),
    "Analyze": ("analyze", "compare", "differentiate", "examine", "organize"),
    "Evaluate": ("evaluate", "justify", "critique", "defend", "assess"),
    "Create": ("design", "develop", "construct", "formulate", "propose"),
}


def _json_list(value):
    try:
        result = json.loads(value or "[]")
        return result if isinstance(result, list) else []
    except (TypeError, ValueError):
        return []


def _subject_report(s):
    subject = {
        "id": s.id, "name": s.name, "code": s.code, "department": s.department,
        "semester": s.semester, "course_outcomes": s.outcomes,
        "syllabus": s.syllabus_text or s.units, "units": s.units, "calendar": s.calendar,
        "resources": s.resources, "website_overview": s.website_overview,
        "announcements": s.announcements,
        "experiential_learning": {"topic": s.case_topic, "scenario": s.case_scenario, "questions": s.case_questions},
        "interactive_prompt": s.interactive_prompt,
        "assessment": {"title": s.exam_title, "marks": s.exam_marks, "questions": s.questions, "answer_key": s.answer_key},
        "obe_improvement_plan": s.improvement_plan, "program_outcomes": s.program_outcomes,
        "co_po_mapping": s.co_po_mapping, "bloom_questions": _json_list(s.bloom_questions),
        "responsible_ai_policy": s.policy,
    }
    return {
        "subject": subject,
        "attendance": [{"student_id": r.student_id, "date": r.attendance_date.isoformat(), "status": r.status}
                       for r in Attendance.query.filter_by(subject_id=s.id).order_by(Attendance.attendance_date).all()],
        "marks": [{"student_id": r.student_id, "assessment": r.assessment, "obtained": r.obtained,
                   "maximum": r.maximum, "percentage": round(r.obtained/r.maximum*100, 2) if r.maximum else None}
                  for r in Mark.query.filter_by(subject_id=s.id).order_by(Mark.id).all()],
        "research_papers": [{"title": r.title, "authors": r.authors, "year": r.year, "doi_url": r.doi_url,
                             "summary": r.summary, "verified": bool(r.verified)}
                            for r in ResearchPaper.query.filter_by(subject_id=s.id).order_by(ResearchPaper.id).all()],
        "quiz": [{"question": q.question, "answer": q.answer}
                 for q in QuizQuestion.query.filter_by(subject_id=s.id).order_by(QuizQuestion.id).all()],
    }


@app.route("/bloom", methods=["GET", "POST"])
def bloom():
    s = get_subject()
    if not s:
        flash("Create a subject first.", "error")
        return redirect(url_for("subject_new"))
    questions = _json_list(s.bloom_questions)
    if request.method == "POST":
        if request.form.get("action") == "save":
            try:
                edited = json.loads(request.form.get("questions_json", "[]"))
                if not isinstance(edited, list):
                    raise ValueError("Expected a list")
                s.bloom_questions = json.dumps(edited, ensure_ascii=False, indent=2)
                db.session.commit()
                flash("Bloom question bank saved for this subject.", "success")
            except (ValueError, TypeError):
                flash("Question data was invalid; no changes were saved.", "error")
            return redirect(url_for("bloom", subject_id=s.id))
        action = request.form.get("action", "generate")
        level = request.form.get("level", "Understand")
        if level not in BLOOM_LEVELS:
            level = "Understand"
        topic = request.form.get("topic", "").strip() or (s.syllabus_text or s.units or s.name).splitlines()[0].strip()
        count = max(1, min(request.form.get("count", 5, type=int) or 5, 20))
        marks = max(1, min(request.form.get("marks", 5, type=int) or 5, 100))
        co = request.form.get("co", "CO1").strip()
        if co not in {f"CO{i}" for i in range(1, 6)}:
            co = "CO1"
        if action == "manual":
            manual_question = request.form.get("manual_question", "").strip()
            if not manual_question:
                flash("Please type a question before adding it.", "error")
                return redirect(url_for("bloom", subject_id=s.id))
            questions.append({"question": manual_question, "topic": topic, "bloom_level": level,
                              "marks": marks, "co": co, "po": "", "answer_outline": "", "manually_added": True})
            s.bloom_questions = json.dumps(questions, ensure_ascii=False, indent=2)
            db.session.commit()
            flash("Your manually written question was added to the saved question bank.", "success")
            return redirect(url_for("bloom", subject_id=s.id))
        templates = {
            "Remember": "Define or list the key terms and principles related to {topic}.",
            "Understand": "Explain {topic} in your own words and illustrate it with an example.",
            "Apply": "Apply the relevant concepts of {topic} to solve a practical problem. Show your steps.",
            "Analyze": "Analyze {topic} by identifying its components, relationships, assumptions, and limitations.",
            "Evaluate": "Evaluate an approach or solution related to {topic}. Justify your judgement using clear criteria.",
            "Create": "Design or develop a solution, model, or plan that addresses a realistic challenge involving {topic}. Explain your choices.",
        }
        verbs = BLOOM_LEVELS[level]
        variations = [
            f"{verbs[1 % len(verbs)].capitalize()} the main concepts involved in {{topic}}, using a different example.",
            f"Given a realistic scenario involving {{topic}}, {verbs[2 % len(verbs)]} an appropriate response and justify it.",
            f"Discuss the assumptions, evidence, or criteria needed to address a problem involving {{topic}}.",
            f"How would you {verbs[0]} the importance of {{topic}} in a practical context?",
        ]
        for i in range(count):
            prompt = templates[level].format(topic=topic) if i == 0 else variations[(i-1) % len(variations)].format(topic=topic)
            questions.append({"question": prompt, "topic": topic, "bloom_level": level, "marks": marks,
                              "co": co, "po": "", "answer_outline": ""})
        s.bloom_questions = json.dumps(questions, ensure_ascii=False, indent=2)
        db.session.commit()
        flash(f"Added {count} editable {level}-level question(s). Review before using in an assessment.", "success")
        return redirect(url_for("bloom", subject_id=s.id))
    return render_template("bloom.html", s=s, questions=questions, levels=list(BLOOM_LEVELS.keys()))


@app.route("/mapping", methods=["GET", "POST"])
def mapping():
    s = get_subject()
    if not s:
        flash("Create a subject first.", "error")
        return redirect(url_for("subject_new"))
    po_names = ["Engineering Knowledge", "Problem Analysis", "Design/Development of Solutions", "Conduct Investigations of Complex Problems", "Engineering Tool Usage", "The Engineer and The World", "Ethics", "Individual and Collaborative Team Work", "Communication", "Project Management and Finance", "Life-Long Learning"]
    if request.method == "POST":
        s.outcomes = request.form.get("outcomes", "")
        rows = ["CO,PO1,PO2,PO3,PO4,PO5,PO6,PO7,PO8,PO9,PO10,PO11"]
        for ci in range(1, 6):
            vals = [request.form.get(f"co{ci}_po{pi}", "-") for pi in range(1, 12)]
            vals = [v if v in ("3", "2", "1") else "-" for v in vals]
            rows.append("CO%d,%s" % (ci, ",".join(vals)))
        s.program_outcomes = "\n".join(f"PO{i}: {name}" for i, name in enumerate(po_names, 1))
        s.co_po_mapping = "\n".join(rows)
        db.session.commit()
        flash("CO1–CO5 and common Engineering PO1–PO11 mapping saved.", "success")
        return redirect(url_for("mapping", subject_id=s.id))
    values = {}
    lines = (s.co_po_mapping or "").splitlines()
    for line in lines[1:]:
        cells = [x.strip() for x in line.split(",")]
        if len(cells) >= 12 and cells[0].startswith("CO"):
            for pi, val in enumerate(cells[1:12], 1): values[f"{cells[0].lower()}_po{pi}"] = val if val in ("3", "2", "1") else "-"
    return render_template("mapping.html", s=s, po_names=po_names, values=values)


@app.route("/bloom/download/<file_format>")
def download_bloom_questions(file_format):
    s = get_subject()
    if not s: return "Create a subject first.", 400
    questions = _json_list(s.bloom_questions)
    if not questions: return "No Bloom questions saved for this subject yet.", 400
    from io import BytesIO
    from flask import send_file
    rows = []
    for i, q in enumerate(questions, 1):
        rows.append((str(i), str(q.get("question", "")), str(q.get("bloom_level", "")), str(q.get("co", "") or "CO1")))
    title = f"{s.name} ({s.code}) — Question Paper"
    if file_format == "doc":
        from docx import Document
        doc = Document(); doc.add_heading(s.name + " (" + s.code + ")", 0)
        doc.add_paragraph("ANSWER THE ALL QUESTIONS")
        table = doc.add_table(rows=1, cols=4); table.style = "Table Grid"
        for cell, txt in zip(table.rows[0].cells, ["Q.No", "QUESTION", "BT LEVEL", "COS"]): cell.text = txt
        for qno, question, level, co in rows:
            cells = table.add_row().cells
            for cell, txt in zip(cells, [qno + ".", question, level, co]): cell.text = txt
        b=BytesIO(); doc.save(b); b.seek(0)
        return send_file(b, as_attachment=True, download_name=f"{s.code}_bloom_questions.docx", mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    if file_format == "pdf":
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.units import mm
        b=BytesIO(); doc=SimpleDocTemplate(b,pagesize=landscape(A4),rightMargin=12*mm,leftMargin=12*mm,topMargin=12*mm,bottomMargin=12*mm)
        styles=getSampleStyleSheet(); story=[Paragraph(title,styles['Title']),Paragraph("ANSWER THE ALL QUESTIONS",styles['Heading2']),Spacer(1,8)]
        data=[["Q.No","QUESTION","BT LEVEL","COS"]]
        from xml.sax.saxutils import escape
        for qno,question,level,co in rows: data.append([qno+".",Paragraph(escape(question),styles['BodyText']),level,co])
        t=Table(data,colWidths=[18*mm,190*mm,28*mm,20*mm],repeatRows=1)
        t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e8eefb')),('GRID',(0,0),(-1,-1),0.5,colors.HexColor('#8ba0c7')),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('PADDING',(0,0),(-1,-1),7)]))
        story.append(t); doc.build(story); b.seek(0)
        return send_file(b, as_attachment=True, download_name=f"{s.code}_bloom_questions.pdf", mimetype="application/pdf")
    if file_format == "jpg":
        from PIL import Image, ImageDraw, ImageFont
        font_path="C:/Windows/Fonts/arial.ttf"
        try: font=ImageFont.truetype(font_path, 20); small=ImageFont.truetype(font_path, 16); head=ImageFont.truetype(font_path, 24)
        except Exception: font=small=head=ImageFont.load_default()
        width=1400; wrap=88; lines=[]
        for qno,question,level,co in rows:
            words=question.split(); cur=""
            for word in words:
                if len(cur)+len(word)+1>wrap: lines.append((cur,font)); cur=word
                else: cur=(cur+" "+word).strip()
            if cur: lines.append((cur,font))
            lines.append((f"BT LEVEL: {level}     COS: {co}",small)); lines.append(("",small))
        height=150+sum(32 if text else 18 for text,f in lines)
        im=Image.new('RGB',(width,height),'white'); d=ImageDraw.Draw(im); d.text((40,25),title,font=head,fill='#153568'); d.text((40,70),'ANSWER THE ALL QUESTIONS',font=font,fill='#111827')
        y=120
        for text,f in lines: d.text((40,y),text,font=f,fill='#111827'); y+=32 if text else 18
        b=BytesIO(); im.save(b,format='JPEG',quality=92); b.seek(0)
        return send_file(b,as_attachment=True,download_name=f"{s.code}_bloom_questions.jpg",mimetype='image/jpeg')
    return "Supported formats: doc, pdf, jpg", 400


@app.route("/reports")
def reports():
    s = get_subject()
    if not s:
        flash("Create a subject first.", "error")
        return redirect(url_for("subject_new"))
    return render_template("reports.html", s=s)


REPORT_PARTS = {
    1: ("Course Planning", ["department", "semester", "course_outcomes", "syllabus", "units", "calendar", "resources"]),
    2: ("Course Website", ["website_overview", "announcements"]),
    3: ("Experiential Learning", ["experiential_learning"]),
    4: ("Interactive Learning", ["interactive_prompt"]),
    5: ("Assessment Design", ["assessment", "bloom_questions"]),
    6: ("Attendance & Marks", ["attendance", "marks"]),
    7: ("OBE Analysis and CO–PO Mapping", ["course_outcomes", "program_outcomes", "co_po_mapping", "marks", "obe_improvement_plan"]),
    8: ("AI Course Agent", ["course_outcomes", "syllabus", "resources"]),
    9: ("Research Integration", ["research_papers"]),
    10: ("Responsible AI", ["responsible_ai_policy"]),
}


def _report_html(s, part, data):
    title, keys = REPORT_PARTS[part]
    blocks = []
    for key in keys:
        value = data.get("subject", {}).get(key, data.get(key, ""))
        value = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, indent=2)
        blocks.append("<h2>" + key.replace("_", " ").title() + "</h2><pre>" + value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;") + "</pre>")
    return ("<!doctype html><html lang='en'><meta charset='utf-8'><title>Part " + str(part) + " - " + title + "</title>"
            "<style>body{font:16px/1.5 Arial,sans-serif;max-width:900px;margin:40px auto;padding:0 20px;color:#172033}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f6f8;padding:14px;border-radius:6px}@media print{body{margin:0}}</style>"
            f"<h1>Part {part} · {title}</h1><p><b>Subject:</b> {s.name} ({s.code}) · {s.semester}</p>" + "".join(blocks) +
            "<footer>Generated by EduOutcome AI. Review content before official submission.</footer></html>")


@app.route("/reports/download/<int:part>")
def download_part_report(part):
    s = get_subject()
    if not s:
        return "Create a subject first.", 400
    if part not in REPORT_PARTS:
        return "Unknown report part.", 404
    response = make_response(_report_html(s, part, _subject_report(s)))
    response.headers["Content-Type"] = "text/html; charset=utf-8"
    response.headers["Content-Disposition"] = f'attachment; filename="{s.code}_part_{part}_report.html"'
    return response


@app.route("/reports/download/all")
def download_all_reports():
    s = get_subject()
    if not s:
        return "Create a subject first.", 400
    from io import BytesIO
    memory_file = BytesIO()
    report = _subject_report(s)
    with zipfile.ZipFile(memory_file, mode="w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(f"{s.code}_complete_course_report.json", json.dumps(report, ensure_ascii=False, indent=2, default=str))
        for part in range(1, 11):
            archive.writestr(f"{s.code}_part_{part}_report.html", _report_html(s, part, report))
        archive.writestr(f"{s.code}_README.txt", "EduOutcome AI report bundle\nIncludes all 10 part reports as HTML plus a complete JSON export.\nOpen an HTML report in a browser and use Print > Save as PDF if needed.\n")
    memory_file.seek(0)
    response = make_response(memory_file.getvalue())
    response.headers["Content-Type"] = "application/zip"
    response.headers["Content-Disposition"] = f'attachment; filename="{s.code}_all_10_part_reports.zip"'
    return response


@app.route("/reports/download/json")
def download_report_json():
    s = get_subject()
    if not s:
        return "Create a subject first.", 400
    return app.response_class(json.dumps(_subject_report(s), ensure_ascii=False, indent=2, default=str),
                              mimetype="application/json",
                              headers={"Content-Disposition": f'attachment; filename="{s.code}_complete_report.json"'})

if __name__ == "__main__":
    app.run(debug=True)
